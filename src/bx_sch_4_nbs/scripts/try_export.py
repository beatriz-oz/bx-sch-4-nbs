from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, select

from bx_sch_4_nbs.config import settings
from bx_sch_4_nbs.database.helpers import engine, get_session
from bx_sch_4_nbs.database.models import User
from bx_sch_4_nbs.database.queries.schedule import list_appointments
from bx_sch_4_nbs.database.types import AppointmentStatus, UserRole
from bx_sch_4_nbs.helpers.authentication import create_access_token
from bx_sch_4_nbs.helpers.common import to_studio_time, utc_now
from bx_sch_4_nbs.helpers.export import received
from bx_sch_4_nbs.main import app

connection = engine.connect()
transaction = connection.begin()
session = Session(bind=connection, join_transaction_mode="create_savepoint")


def override_session():
    try:
        yield session
        session.commit()
    except Exception:
        session.rollback()
        raise


app.dependency_overrides[get_session] = override_session
client = TestClient(app, raise_server_exceptions=False)


def check(label: str, ok: bool, detail: object = "") -> None:
    print("OK    " if ok else "ERRO  ", label, detail if not ok else "")


admin = session.exec(select(User).where(User.role == UserRole.SUPER_ADMIN)).first()
regular = session.exec(select(User).where(User.role == UserRole.USER)).first()
assert admin is not None and regular is not None
ADMIN = {"Authorization": f"Bearer {create_access_token(admin)}"}
CLIENT = {"Authorization": f"Bearer {create_access_token(regular)}"}
URL = "/v1/admin/appointments/export"

# The last 3 months (the demo data has completed appointments, no-shows and cancellations there).
today = to_studio_time(utc_now()).date()
start, end = today - timedelta(days=95), today
expected_rows = list_appointments(session, start, end)
assert expected_rows, "Run seed_demo first: there are no appointments in the last 3 months."

# --- Access and validation ---
check("no login -> 401", client.get(URL, params={"start_date": start, "end_date": end}).status_code == 401)
check(
    "a client -> 403", client.get(URL, headers=CLIENT, params={"start_date": start, "end_date": end}).status_code == 403
)
r = client.get(URL, headers=ADMIN, params={"start_date": end, "end_date": start})
check("end before start -> 422", r.status_code == 422, r.status_code)
r = client.get(URL, headers=ADMIN, params={"start_date": "2026-01-01", "end_date": "2028-06-01"})
check("more than 2 years -> 422", r.status_code == 422, r.status_code)

# --- The file ---
r = client.get(URL, headers=ADMIN, params={"start_date": start, "end_date": end})
check("export -> 200", r.status_code == 200, r.text[:200])
check(
    "an .xlsx file", r.headers["content-type"].startswith("application/vnd.openxmlformats"), r.headers["content-type"]
)
check(
    "downloaded with a clear name",
    f'filename="marcacoes_{start.isoformat()}_{end.isoformat()}.xlsx"' in r.headers["content-disposition"],
    r.headers["content-disposition"],
)

sheet = load_workbook(BytesIO(r.content)).active
assert sheet is not None
header = [cell.value for cell in sheet[1]]
check("Portuguese headers by default", header[:3] == ["Data", "Hora", "Cliente"] and "Recebido" in header, header)
data_rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if isinstance(row[0], datetime | date)]
check("one row per appointment", len(data_rows) == len(expected_rows), (len(data_rows), len(expected_rows)))

first = data_rows[0]
check("dates are real dates", isinstance(first[0], datetime | date), type(first[0]))
check("times are real times", isinstance(first[1], time), type(first[1]))
col = {name: header.index(name) for name in header}
check("prices are numbers", isinstance(first[col["Preço estimado"]], int | float), type(first[col["Preço estimado"]]))

expected_received = sum((received(a, settings.deposit_amount) for a, _ in expected_rows), Decimal(0))
in_file = sum(Decimal(str(row[col["Recebido"]])) for row in data_rows)
check("received = completed prices + kept deposits", in_file == expected_received, (in_file, expected_received))

total_row = [row for row in sheet.iter_rows(values_only=True) if row[0] == "Total"]
check(
    "a total row with a SUM formula",
    bool(total_row) and str(total_row[0][col["Recebido"]]).startswith("=SUM("),
    total_row,
)

statuses = {row[col["Estado"]] for row in data_rows}
check("statuses in Portuguese", statuses <= {"Marcada", "Confirmada", "Concluída", "Cancelada", "Faltou"}, statuses)
completed = [a for a, _ in expected_rows if a.status == AppointmentStatus.COMPLETED]
check("the period has completed appointments", bool(completed))
check("the header row stays visible", sheet.freeze_panes == "A2", sheet.freeze_panes)

# --- English ---
r = client.get(URL, headers=ADMIN, params={"start_date": start, "end_date": end, "language": "en"})
sheet_en = load_workbook(BytesIO(r.content)).active
assert sheet_en is not None
check("English when asked", [c.value for c in sheet_en[1]][:3] == ["Date", "Time", "Client"])
check("English file name", "appointments_" in r.headers["content-disposition"], r.headers["content-disposition"])

# --- An empty period still gives a valid file ---
r = client.get(URL, headers=ADMIN, params={"start_date": "2099-01-01", "end_date": "2099-01-31"})
empty = load_workbook(BytesIO(r.content)).active
check("empty period -> header and total only", r.status_code == 200 and empty is not None and empty.max_row <= 3)

session.close()
transaction.rollback()
connection.close()
